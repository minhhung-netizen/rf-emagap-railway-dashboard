import json
import re
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from app.database import SignalStore
from app.services import trade_reconciliation as service


HEADERS = ["STT", "Giao dịch", "Mã", "Khối lượng khớp", "Giá khớp bình quân", "Giá trị khớp"]


def workbook(rows=None, day="05/01/2026", account="SYNTH0001", subaccount="Thường", totals=None):
    """Synthetic report only; no user statement data or account identifiers."""
    book = Workbook()
    sheet = book.active
    sheet.title = "match_orders"
    sheet.append([f"Ngày giao dịch: {day} 15:30:00"])
    sheet.append([f"Số tài khoản: {account}"])
    sheet.append([f"Tiểu khoản: {subaccount}"])
    if totals is not None:
        for side, amount in totals.items():
            sheet.append([f"Tổng giá trị {side}: {amount}"])
    sheet.append(HEADERS)
    for i, row in enumerate(rows or [], 1):
        sheet.append([i, *row])
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def parse(rows, **kwargs):
    return service.parse_workbook(workbook(rows, **kwargs), "report.xlsx")


def buy(qty=100, price=10000, ticker="ABC"):
    return ["Mua", ticker, qty, price, qty * price]


def sell(qty=50, price=12000, ticker="ABC"):
    return ["Bán", ticker, qty, price, qty * price]


class WorkbookTests(unittest.TestCase):
    def test_template_numeric_strings_totals_and_authoritative_average(self):
        report = parse([["Mua", "abc", "1,000", "11,487.500000", "11,487,500"],
                        ["Bán", "ABC", "700", "11,542.857143", "8,080,000"]],
                       totals={"mua": "11,487,500", "bán": "8,080,000", "khớp": "19,567,500"})
        self.assertEqual(report["trade_date"], "2026-01-05")
        self.assertEqual(report["currency"], "VND")
        self.assertEqual(len(report["rows"]), 1)
        row = report["rows"][0]
        self.assertEqual(row["residual_buy_quantity"], 300)
        self.assertEqual(row["matched_quantity"], 700)
        self.assertEqual(row["buy_average_price"], 11487.5)
        self.assertAlmostEqual(row["sell_average_price"], 8080000 / 700)
        self.assertEqual(row["buy_amount_exact"], "11487500")

    def test_duplicate_rows_aggregate_and_digest_ignores_filename_order_numeric_format(self):
        original = service.parse_workbook(workbook([buy(50), buy(50), sell(20)]), "one.xlsx")
        reordered = service.parse_workbook(workbook([sell(20), ["Mua", "ABC", "100.0", "10000.000000", "1000000.00"]]), "two.xlsx")
        self.assertEqual(original["digest"], reordered["digest"])
        self.assertEqual(original["rows"], reordered["rows"])
        self.assertEqual(original["rows"][0]["buy_quantity"], 100)

    def test_masked_label_and_account_subaccount_scopes_are_isolated(self):
        a = parse([buy()], account="SYNTH0001")
        b = parse([buy()], account="OTHER0001")
        c = parse([buy()], account="SYNTH0001", subaccount="Ký quỹ")
        self.assertNotIn("SYNTH0001", json.dumps(a, ensure_ascii=False))
        self.assertIn("0001", a["account_label"])
        self.assertNotEqual(a["scope"], b["scope"])
        self.assertNotEqual(a["scope"], c["scope"])

    def test_incorrect_A1_dimension_does_not_truncate_rows(self):
        data = workbook([buy(), sell()])
        target = BytesIO()
        with zipfile.ZipFile(BytesIO(data)) as source, zipfile.ZipFile(target, "w") as destination:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', content)
                destination.writestr(item, content)
        result = service.parse_workbook(target.getvalue(), "dimension.xlsx")
        self.assertEqual(result["rows"][0]["sell_quantity"], 50)

    def test_reject_missing_invalid_future_date_and_missing_account(self):
        for values in [dict(day="31/02/2026"), dict(day="05/01/2099"), dict(account="")]:
            with self.subTest(values=values), self.assertRaises(ValueError):
                parse([buy()], **values)
        book = Workbook()
        book.active.append(HEADERS)
        book.active.append([1, *buy()])
        buffer = BytesIO(); book.save(buffer); book.close()
        with self.assertRaisesRegex(ValueError, "Ngày giao dịch"):
            service.parse_workbook(buffer.getvalue(), "missing.xlsx")

    def test_reject_invalid_numeric_side_ticker_formulas_and_totals(self):
        invalid = [
            ["Mua", "ABC", "1,00", 10000, 1000000],
            ["Mua", "ABC", 0.5, 10000, 5000],
            ["Mua", "ABC", 0, 10000, 0],
            ["Mua", "ABC", 100, "NaN", 1000000],
            ["Mua", "ABC", 100, "Infinity", 1000000],
            ["Mua", "ABC", 100, "1e-999999999", "1e-999999997"],
            ["Mua", "ABC", 100, 9999, 1000000],
            ["Buy", "ABC", 100, 10000, 1000000],
            ["Mua", "<script>", 100, 10000, 1000000],
            ["Mua", "ABC", "=50+50", 10000, 1000000],
        ]
        for row in invalid:
            with self.subTest(row=row), self.assertRaises(ValueError):
                parse([row])
        with self.assertRaisesRegex(ValueError, "Tổng giá trị"):
            parse([buy()], totals={"mua": 999999})

    def test_reject_malformed_empty_multiple_report_sheets_and_limits(self):
        for data, name in [(b"not an excel", "report.xlsx"), (workbook([buy()]), "report.xls"), (workbook([]), "empty.xlsx")]:
            with self.subTest(name=name), self.assertRaises(ValueError):
                service.parse_workbook(data, name)
        book = Workbook()
        for sheet in (book.active, book.create_sheet("second")):
            sheet.append(["Ngày giao dịch: 05/01/2026"])
            sheet.append(["Số tài khoản: SYNTH0001"])
            sheet.append(HEADERS)
            sheet.append([1, *buy()])
        buffer = BytesIO(); book.save(buffer); book.close()
        with self.assertRaisesRegex(ValueError, "đúng một bảng"):
            service.parse_workbook(buffer.getvalue(), "multiple.xlsx")
        with patch.object(service, "MAX_FILE_BYTES", 1), self.assertRaisesRegex(ValueError, "10 MB"):
            parse([buy()])
        with patch.object(service, "MAX_EXPANDED_BYTES", 1), self.assertRaisesRegex(ValueError, "50 MB"):
            parse([buy()])


class ReconciliationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = SignalStore(Path(self.temp.name) / "test.db")

    def tearDown(self):
        self.temp.cleanup()

    def add(self, rows, *, replace=False, **kwargs):
        parsed = parse(rows, **kwargs)
        return service.import_report(self.store, parsed, "tester", replace), parsed

    def test_daily_net_residual_and_weighted_buy_cost_then_partial_sell(self):
        self.add([buy(100, 10000), sell(40)])
        self.add([buy(100, 20000), sell(50)], day="06/01/2026")
        report = service.build_report(self.store)
        self.assertEqual(report["method"], "daily_net_inventory_with_base")
        position = report["positions"][0]
        self.assertEqual(position["quantity"], 110)
        self.assertEqual(position["total_cost"], 1600000)
        self.assertAlmostEqual(position["average_cost"], 1600000 / 110)
        newest = report["rows"][0]
        self.assertEqual(newest["opening_quantity"], 60)
        self.assertEqual(newest["opening_average_cost"], 10000)
        self.assertEqual(newest["residual_buy_cost"], 1000000)
        self.assertEqual(newest["closing_quantity"], 110)
        self.add([buy(20, 25000), sell(100)], day="07/01/2026")
        report = service.build_report(self.store)
        position = report["positions"][0]
        self.assertEqual(position["quantity"], 30)
        self.assertAlmostEqual(position["average_cost"], 1600000 / 110)
        self.assertAlmostEqual(position["total_cost"], 30 * 1600000 / 110)
        newest = report["rows"][0]
        self.assertEqual(newest["matched_quantity"], 20)
        self.assertEqual(newest["tracked_sell_quantity"], 80)
        self.assertEqual(newest["ignored_sell_quantity"], 0)
        self.assertEqual(newest["residual_buy_cost"], 0)
        self.add([buy(100, 25000)], day="08/01/2026")
        position = service.build_report(self.store)["positions"][0]
        self.assertEqual(position["quantity"], 130)
        self.assertAlmostEqual(position["average_cost"], (30 * 1600000 / 110 + 2500000) / 130)

    def test_equal_has_no_effect_excess_sales_close_cost_and_new_buy_resets_average(self):
        self.add([buy(100, 10000)])
        self.add([buy(100, 20000), sell(100)], day="06/01/2026")
        self.add([buy(50, 20000), sell(1000)], day="07/01/2026")
        self.add([sell(99999)], day="08/01/2026")
        report = service.build_report(self.store)
        self.assertEqual(len(report["rows"]), 4)
        self.assertEqual(report["positions"][0]["quantity"], 0)
        self.assertIsNone(report["positions"][0]["average_cost"])
        self.assertEqual(report["positions"][0]["total_cost"], 0)
        self.assertEqual(report["positions"][0]["status"], "closed")
        self.assertEqual(report["totals"]["outstanding_tickers"], 0)
        self.assertEqual(report["totals"]["ignored_rows_count"], 2)
        self.assertEqual(report["totals"]["ignored_sell_quantity"], 850 + 99999)
        close = next(r for r in report["rows"] if r["trade_date"] == "2026-01-07")
        self.assertEqual(close["tracked_sell_quantity"], 100)
        self.assertEqual(close["ignored_sell_quantity"], 850)
        self.assertIsNone(close["closing_average_cost"])
        self.assertNotIn("unmatched_sell_quantity", report["totals"])
        self.add([buy(100, 20000)], day="09/01/2026")
        position = service.build_report(self.store)["positions"][0]
        self.assertEqual(position["quantity"], 100)
        self.assertEqual(position["average_cost"], 20000)

    def test_sell_only_import_does_not_offset_future_buy(self):
        self.add([sell(1000)])
        report = service.build_report(self.store)
        self.assertEqual(report["positions"][0]["quantity"], 0)
        self.assertEqual(report["rows"][0]["ignored_sell_quantity"], 1000)
        self.add([buy(10)], day="06/01/2026")
        self.assertEqual(service.build_report(self.store)["positions"][0]["quantity"], 10)

    def test_out_of_order_and_backdated_replacement_recompute_all_later_balances(self):
        self.add([buy(100, 20000)], day="07/01/2026")
        self.add([buy(100, 10000)], day="05/01/2026")
        report = service.build_report(self.store)
        self.assertEqual(report["positions"][0]["average_cost"], 15000)
        self.assertEqual(report["rows"][0]["opening_quantity"], 100)
        self.add([buy(50, 10000)], day="05/01/2026", replace=True)
        report = service.build_report(self.store)
        self.assertEqual(report["rows"][0]["opening_quantity"], 50)
        self.assertEqual(report["positions"][0]["quantity"], 150)
        self.assertAlmostEqual(report["positions"][0]["average_cost"], 2500000 / 150)

    def test_duplicate_conflict_and_atomic_replacement(self):
        result, parsed = self.add([buy()])
        self.assertEqual(result["inserted"], 1)
        self.assertEqual(service.assess_import(self.store, parsed), "duplicate")
        self.assertTrue(service.import_report(self.store, parsed, "another")["duplicate"])
        changed = parse([buy(200)])
        self.assertEqual(service.assess_import(self.store, changed), "conflict")
        before = service.build_report(self.store)
        with self.assertRaises(service.ReconciliationConflict):
            service.import_report(self.store, changed, "another")
        self.assertEqual(service.build_report(self.store), before)
        result = service.import_report(self.store, changed, "another", True)
        self.assertTrue(result["replaced"])
        self.assertEqual(service.build_report(self.store)["positions"][0]["quantity"], 200)
        self.assertEqual(service.build_report(self.store)["totals"]["imports_count"], 1)

    def test_accounts_and_subaccounts_never_mix_costs(self):
        _, first = self.add([buy(100, 10000)], account="SYNTH0001")
        self.add([buy(100, 20000)], account="OTHER0001")
        self.add([buy(100, 30000)], account="SYNTH0001", subaccount="Ký quỹ")
        report = service.build_report(self.store)
        self.assertEqual(len(report["scopes"]), 3)
        self.assertEqual(len(report["positions"]), 1)
        self.assertTrue(report["selected_scope"])
        scoped = service.build_report(self.store, first["scope"])
        self.assertEqual(len(scoped["positions"]), 1)
        self.assertEqual(scoped["positions"][0]["average_cost"], 10000)
        with self.assertRaises(ValueError):
            service.build_report(self.store, "unknown")

    def test_colliding_masked_accounts_have_stable_distinct_selector_labels(self):
        _, first = self.add([buy(100, 10000)], account="SYNTH0001")
        _, second = self.add([buy(100, 20000)], account="OTHER0001")
        self.assertEqual(first["account_label"], second["account_label"])
        report = service.build_report(self.store)
        labels = {item["scope"]: item["account_label"] for item in report["scopes"]}
        self.assertEqual(len(set(labels.values())), 2)
        for parsed in (first, second):
            self.assertEqual(labels[parsed["scope"]], parsed["account_label"] + " · #" + parsed["scope"][-6:])
        selected = report["selected_scope"]
        self.add([buy(10)], day="06/01/2026", account="SYNTH0001")
        later = service.build_report(self.store)
        self.assertEqual(later["selected_scope"], selected)
        self.assertEqual({item["scope"]: item["account_label"] for item in later["scopes"]}, labels)

    def test_notes_survive_buy_to_sell_replacement_and_removed_tickers_are_pruned(self):
        _, parsed = self.add([buy(), buy(10, ticker="XYZ")])
        saved = service.save_note(self.store, parsed["scope"], parsed["trade_date"], "abc", " Theo dõi mua ", "tester")
        self.assertEqual(saved["note"], "Theo dõi mua")
        report = service.build_report(self.store)
        self.assertEqual(next(r for r in report["rows"] if r["ticker"] == "ABC")["note"], "Theo dõi mua")
        self.add([buy(200)], replace=True)
        self.assertEqual(service.build_report(self.store)["rows"][0]["note"], "Theo dõi mua")
        self.add([sell()], replace=True)
        with self.store.connect() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM reconciliation_notes").fetchone()[0], 1)
        self.assertEqual(service.build_report(self.store)["rows"][0]["note"], "Theo dõi mua")
        service.save_note(self.store, parsed["scope"], parsed["trade_date"], "ABC", "Bán phần dư", "tester")
        self.assertEqual(service.build_report(self.store)["rows"][0]["note"], "Bán phần dư")
        self.add([buy(), sell(100)], replace=True)
        service.save_note(self.store, parsed["scope"], parsed["trade_date"], "ABC", "Đối ứng đủ", "tester")
        self.assertEqual(service.build_report(self.store)["rows"][0]["note"], "Đối ứng đủ")
        with self.assertRaises(ValueError):
            service.save_note(self.store, parsed["scope"], parsed["trade_date"], "XYZ", "removed ticker", "tester")
        with self.assertRaises(ValueError):
            service.save_note(self.store, parsed["scope"], parsed["trade_date"], "XYZ", "x" * 2001, "tester")

    def test_base_is_already_owned_and_surplus_sales_never_consume_it(self):
        _, parsed = self.add([buy(2000, 20000, "CII")])
        saved = service.save_base(self.store, parsed["scope"], "cii", 10000, "tester")
        self.assertEqual(saved["base_quantity"], 10000)
        report = service.build_report(self.store)
        position = report["positions"][0]
        self.assertEqual(position["base_quantity"], 10000)
        self.assertEqual(position["quantity"], 2000)
        self.assertEqual(position["total_quantity"], 12000)
        self.assertEqual(position["average_cost"], 20000)
        self.assertEqual(position["total_cost"], 40000000)
        self.assertEqual(report["bases"][0]["ticker"], "CII")
        row = report["rows"][0]
        self.assertEqual(row["opening_total_quantity"], 10000)
        self.assertEqual(row["closing_total_quantity"], 12000)
        self.add([sell(500, ticker="CII")], day="06/01/2026")
        report = service.build_report(self.store)
        self.assertEqual(report["positions"][0]["total_quantity"], 11500)
        self.assertEqual(report["positions"][0]["average_cost"], 20000)
        self.add([sell(9000, ticker="CII")], day="07/01/2026")
        report = service.build_report(self.store)
        position = report["positions"][0]
        self.assertEqual(position["quantity"], 0)
        self.assertEqual(position["base_quantity"], 10000)
        self.assertEqual(position["total_quantity"], 10000)
        self.assertIsNone(position["average_cost"])
        self.assertEqual(position["total_cost"], 0)
        self.assertEqual(position["status"], "base_only")
        self.assertEqual(report["rows"][0]["tracked_sell_quantity"], 1500)
        self.assertEqual(report["rows"][0]["ignored_sell_quantity"], 7500)
        self.assertEqual(report["rows"][0]["closing_total_quantity"], 10000)
        self.assertEqual(report["totals"]["base_quantity"], 10000)
        self.assertEqual(report["totals"]["total_quantity"], 10000)

    def test_base_can_have_no_source_trade_and_zero_clears_without_changing_cost(self):
        _, parsed = self.add([buy(100, 10000)])
        service.save_base(self.store, parsed["scope"], "CII", 10000, "tester")
        report = service.build_report(self.store)
        by_ticker = {p["ticker"]: p for p in report["positions"]}
        self.assertEqual(by_ticker["CII"]["total_quantity"], 10000)
        self.assertIsNone(by_ticker["CII"]["first_trade_date"])
        self.assertIsNone(by_ticker["CII"]["average_cost"])
        self.assertEqual(by_ticker["CII"]["status"], "base_only")
        self.assertEqual(report["totals"]["outstanding_tickers"], 1)
        self.assertEqual(report["totals"]["outstanding_quantity"], 100)
        service.save_base(self.store, parsed["scope"], "ABC", 5000, "tester")
        before = service.build_report(self.store)["positions"][0]
        self.assertEqual(before["total_quantity"], 5100)
        service.save_base(self.store, parsed["scope"], "ABC", 0, "tester")
        after = service.build_report(self.store)["positions"][0]
        self.assertEqual(after["quantity"], before["quantity"])
        self.assertEqual(after["average_cost"], before["average_cost"])
        self.assertEqual(after["total_cost"], before["total_cost"])
        self.assertEqual(after["total_quantity"], 100)
        service.save_base(self.store, parsed["scope"], "CII", 0, "tester")
        self.assertEqual(service.build_report(self.store)["totals"]["base_quantity"], 0)

    def test_base_quantity_validation_and_unknown_account_are_atomic(self):
        _, parsed = self.add([buy()])
        service.save_base(self.store, parsed["scope"], "ABC", 10000, "tester")
        before = service.build_report(self.store)
        for quantity in (-1, True, False, 0.5, 1.0, "100", None, 10**12 + 1):
            with self.subTest(quantity=quantity), self.assertRaises(ValueError):
                service.save_base(self.store, parsed["scope"], "ABC", quantity, "tester")
        for ticker in ("", "<script>", "a" * 21, None):
            with self.subTest(ticker=ticker), self.assertRaises(ValueError):
                service.save_base(self.store, parsed["scope"], ticker, 100, "tester")
        with self.assertRaises(ValueError):
            service.save_base(self.store, "unknown", "ABC", 100, "tester")
        self.assertEqual(service.build_report(self.store), before)
        service.save_base(self.store, parsed["scope"], "ABC", 10**12, "tester")
        self.assertEqual(service.build_report(self.store)["positions"][0]["base_quantity"], 10**12)

    def test_base_and_sales_are_isolated_by_account_and_subaccount(self):
        _, first = self.add([buy(100, 10000)], account="SYNTH0001")
        _, second = self.add([buy(100, 20000)], account="OTHER0001")
        _, third = self.add([buy(100, 30000)], account="SYNTH0001", subaccount="Ký quỹ")
        service.save_base(self.store, first["scope"], "ABC", 10000, "tester")
        self.add([sell(1000)], day="06/01/2026", account="SYNTH0001")
        report = service.build_report(self.store, first["scope"])
        self.assertEqual(report["positions"][0]["quantity"], 0)
        self.assertEqual(report["positions"][0]["total_quantity"], 10000)
        for parsed, cost in ((second, 20000), (third, 30000)):
            report = service.build_report(self.store, parsed["scope"])
            self.assertEqual(report["positions"][0]["quantity"], 100)
            self.assertEqual(report["positions"][0]["base_quantity"], 0)
            self.assertEqual(report["positions"][0]["average_cost"], cost)
            self.assertEqual(report["bases"], [])

    def test_backdated_sell_and_replacement_replay_inventory_above_current_base(self):
        _, parsed = self.add([buy(100, 20000)], day="07/01/2026")
        service.save_base(self.store, parsed["scope"], "ABC", 10000, "tester")
        self.add([buy(100, 10000)], day="05/01/2026")
        _, sold = self.add([sell(50)], day="06/01/2026")
        service.save_note(self.store, sold["scope"], sold["trade_date"], "ABC", "Bán dư mua", "tester")
        report = service.build_report(self.store)
        self.assertEqual(report["positions"][0]["quantity"], 150)
        self.assertAlmostEqual(report["positions"][0]["average_cost"], 2500000 / 150)
        self.add([buy(20, 10000)], day="05/01/2026", replace=True)
        report = service.build_report(self.store)
        self.assertEqual(report["positions"][0]["quantity"], 100)
        self.assertEqual(report["positions"][0]["average_cost"], 20000)
        self.assertEqual(report["positions"][0]["total_quantity"], 10100)
        self.assertEqual(report["rows"][0]["opening_total_quantity"], 10000)
        sale = next(r for r in report["rows"] if r["trade_date"] == "2026-01-06")
        self.assertEqual(sale["tracked_sell_quantity"], 20)
        self.assertEqual(sale["ignored_sell_quantity"], 30)
        self.assertEqual(sale["note"], "Bán dư mua")

    def test_delete_purchase_and_sale_replay_later_inventory_without_touching_bases_or_other_accounts(self):
        _, first = self.add([buy(100, 10000)])
        _, second = self.add([buy(100, 20000)], day="06/01/2026")
        self.add([sell(50)], day="07/01/2026")
        _, other = self.add([buy(100, 30000)], account="OTHER0001")
        scope = first["scope"]
        service.save_base(self.store, scope, "ABC", 10000, "tester")
        service.save_note(self.store, scope, first["trade_date"], "ABC", "Xóa ngày mua đầu", "tester")
        service.save_note(self.store, scope, second["trade_date"], "ABC", "Giữ ngày mua sau", "tester")
        service.save_note(self.store, other["scope"], other["trade_date"], "ABC", "Tài khoản khác", "tester")
        before_other = service.build_report(self.store, other["scope"])
        report = service.build_report(self.store, scope)
        self.assertEqual(report["positions"][0]["quantity"], 150)
        self.assertEqual(report["positions"][0]["average_cost"], 15000)
        target = next(i for i in report["imports"] if i["trade_date"] == first["trade_date"])
        result = service.delete_report(self.store, scope, target["id"], target["revision"])
        self.assertEqual(result, dict(scope=scope, import_id=target["id"], trade_date=first["trade_date"],
                                      deleted_imports_count=1, deleted_notes_count=1, deleted_bases_count=0))
        report = service.build_report(self.store, scope)
        self.assertEqual(report["positions"][0]["quantity"], 50)
        self.assertEqual(report["positions"][0]["average_cost"], 20000)
        self.assertEqual(report["positions"][0]["total_cost"], 1000000)
        self.assertEqual(report["positions"][0]["base_quantity"], 10000)
        self.assertEqual(report["rows"][1]["note"], "Giữ ngày mua sau")
        after_other = service.build_report(self.store, other["scope"])
        self.assertEqual({k: v for k, v in after_other.items() if k != "scopes"},
                         {k: v for k, v in before_other.items() if k != "scopes"})
        sale = next(i for i in report["imports"] if i["trade_date"] == "2026-01-07")
        service.delete_report(self.store, scope, sale["id"], sale["revision"])
        report = service.build_report(self.store, scope)
        self.assertEqual(report["positions"][0]["quantity"], 100)
        self.assertEqual(report["positions"][0]["total_cost"], 2000000)
        self.assertEqual(report["positions"][0]["total_quantity"], 10100)

    def test_last_report_delete_keeps_selectable_base_only_scope_and_reimport_restores_tracking(self):
        _, parsed = self.add([buy(100, 20000)])
        scope = parsed["scope"]
        service.save_base(self.store, scope, "ABC", 10000, "tester")
        service.save_base(self.store, scope, "XYZ", 0, "tester")
        service.save_note(self.store, scope, parsed["trade_date"], "ABC", "Xóa cùng báo cáo", "tester")
        original = service.build_report(self.store, scope)["imports"][0]
        service.delete_report(self.store, scope, original["id"], original["revision"])
        report = service.build_report(self.store, scope)
        self.assertEqual(report["selected_scope"], scope)
        self.assertEqual(report["scopes"][0], dict(scope=scope, account_label="Tài khoản · #" + scope[-6:],
                                                  first_trade_date=None, last_trade_date=None, imports_count=0))
        self.assertEqual(report["imports"], [])
        self.assertEqual(report["rows"], [])
        self.assertEqual(report["totals"]["total_quantity"], 10000)
        self.assertRegex(report["revision"], r"^[a-f0-9]{64}$")
        service.save_base(self.store, scope, "XYZ", 500, "tester")
        self.assertEqual(service.build_report(self.store, scope)["totals"]["base_quantity"], 10500)
        service.import_report(self.store, parsed, "tester")
        report = service.build_report(self.store, scope)
        self.assertNotEqual(report["imports"][0]["id"], original["id"])
        self.assertNotEqual(report["imports"][0]["revision"], original["revision"])
        self.assertEqual(report["positions"][0]["quantity"], 100)
        self.assertEqual(report["positions"][0]["total_quantity"], 10100)
        self.assertEqual(report["rows"][0]["note"], "")
        self.assertEqual(report["scopes"][0]["account_label"], parsed["account_label"])

    def test_last_report_without_bases_removes_scope_and_empty_report_has_no_revision(self):
        _, parsed = self.add([buy()])
        report = service.build_report(self.store, parsed["scope"])
        target = report["imports"][0]
        service.delete_report(self.store, parsed["scope"], target["id"], target["revision"])
        report = service.build_report(self.store)
        self.assertEqual(report["scopes"], [])
        self.assertIsNone(report["selected_scope"])
        self.assertIsNone(report["revision"])
        self.assertEqual(report["positions"], [])
        with self.assertRaises(ValueError):
            service.build_report(self.store, parsed["scope"])

    def test_clear_scope_deletes_imports_notes_bases_only_for_selected_account(self):
        _, first = self.add([buy()])
        self.add([sell()], day="06/01/2026")
        _, other = self.add([buy(100, 20000)], account="OTHER0002")
        scope = first["scope"]
        service.save_note(self.store, scope, first["trade_date"], "ABC", "Mua", "tester")
        service.save_note(self.store, scope, "2026-01-06", "ABC", "Bán", "tester")
        service.save_base(self.store, scope, "ABC", 10000, "tester")
        service.save_base(self.store, scope, "XYZ", 500, "tester")
        service.save_note(self.store, other["scope"], other["trade_date"], "ABC", "Giữ", "tester")
        service.save_base(self.store, other["scope"], "ABC", 20000, "tester")
        other_before = service.build_report(self.store, other["scope"])
        report = service.build_report(self.store, scope)
        result = service.delete_scope(self.store, scope, report["revision"])
        self.assertEqual(result, dict(scope=scope, deleted_imports_count=2, deleted_notes_count=2, deleted_bases_count=2, deleted_fees_count=0))
        self.assertEqual(service.build_report(self.store, other["scope"]), {
            **other_before, "scopes": [s for s in other_before["scopes"] if s["scope"] == other["scope"]]})
        with self.store.connect() as conn:
            for table in ("reconciliation_imports", "reconciliation_notes", "reconciliation_bases"):
                self.assertEqual(conn.execute(f"SELECT COUNT(*) FROM {table} WHERE scope=?", (scope,)).fetchone()[0], 0)
            for table in ("ledger_entries", "ledger_closes", "nav_snapshots"):
                self.assertEqual(conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0], 0)

    def test_clear_base_only_scope_includes_orphan_notes_and_preserves_other_scopes(self):
        _, first = self.add([buy()])
        _, other = self.add([buy(100, 20000)], account="OTHER0001")
        scope = first["scope"]
        service.save_base(self.store, scope, "ABC", 10000, "tester")
        report = service.build_report(self.store, scope)
        target = report["imports"][0]
        service.delete_report(self.store, scope, target["id"], target["revision"])
        with self.store.connect() as conn:
            conn.execute("INSERT INTO reconciliation_notes(scope,trade_date,ticker,note,updated_by,updated_at) VALUES(?,?,?,?,?,?)",
                         (scope, first["trade_date"], "ABC", "Legacy orphan", "tester", "2026-01-05T00:00:00+00:00"))
        report = service.build_report(self.store, scope)
        result = service.delete_scope(self.store, scope, report["revision"])
        self.assertEqual(result, dict(scope=scope, deleted_imports_count=0, deleted_notes_count=1, deleted_bases_count=1, deleted_fees_count=0))
        self.assertEqual(service.build_report(self.store)["selected_scope"], other["scope"])

    def test_report_delete_rejects_replacement_and_note_edit_with_no_partial_deletion(self):
        _, parsed = self.add([buy()])
        scope = parsed["scope"]
        original = service.build_report(self.store, scope)["imports"][0]
        self.add([buy(200)], replace=True)
        replaced = service.build_report(self.store, scope)
        self.assertEqual(replaced["imports"][0]["id"], original["id"])
        with self.assertRaises(service.ReconciliationConflict):
            service.delete_report(self.store, scope, original["id"], original["revision"])
        self.assertEqual(service.build_report(self.store, scope), replaced)
        stale = replaced["imports"][0]
        service.save_note(self.store, scope, parsed["trade_date"], "ABC", "New note", "tester")
        current = service.build_report(self.store, scope)
        with self.assertRaises(service.ReconciliationConflict):
            service.delete_report(self.store, scope, stale["id"], stale["revision"])
        self.assertEqual(service.build_report(self.store, scope), current)

    def test_scope_delete_rejects_new_import_note_and_base_edits_atomically(self):
        _, parsed = self.add([buy()])
        scope = parsed["scope"]
        changes = (
            lambda: self.add([buy(100, 20000)], day="06/01/2026"),
            lambda: service.save_note(self.store, scope, parsed["trade_date"], "ABC", "New note", "tester"),
            lambda: service.save_base(self.store, scope, "ABC", 10000, "tester"),
        )
        for change in changes:
            stale = service.build_report(self.store, scope)
            change()
            current = service.build_report(self.store, scope)
            with self.assertRaises(service.ReconciliationConflict):
                service.delete_scope(self.store, scope, stale["revision"])
            self.assertEqual(service.build_report(self.store, scope), current)

    def test_revisions_are_deterministic_and_include_same_content_audit_edits(self):
        _, parsed = self.add([buy()])
        scope = parsed["scope"]
        with patch.object(service, "utc_now_iso", return_value="2026-01-05T00:00:00.000001+00:00"):
            service.save_note(self.store, scope, parsed["trade_date"], "ABC", "Same note", "tester")
            service.save_base(self.store, scope, "ABC", 10000, "tester")
        original = service.build_report(self.store, scope)
        self.assertEqual(service.build_report(self.store, scope), original)
        for revision in (original["revision"], original["imports"][0]["revision"]):
            self.assertRegex(revision, r"^[a-f0-9]{64}$")
        with patch.object(service, "utc_now_iso", return_value="2026-01-05T00:00:00.000002+00:00"):
            service.save_note(self.store, scope, parsed["trade_date"], "ABC", "Same note", "tester")
        updated_note = service.build_report(self.store, scope)
        self.assertNotEqual(updated_note["revision"], original["revision"])
        self.assertNotEqual(updated_note["imports"][0]["revision"], original["imports"][0]["revision"])
        with patch.object(service, "utc_now_iso", return_value="2026-01-05T00:00:00.000003+00:00"):
            service.save_base(self.store, scope, "ABC", 10000, "tester")
        updated_base = service.build_report(self.store, scope)
        self.assertNotEqual(updated_base["revision"], updated_note["revision"])
        self.assertEqual(updated_base["imports"][0]["revision"], updated_note["imports"][0]["revision"])

    def test_missing_wrong_scope_and_invalid_revision_never_delete_anything(self):
        _, first = self.add([buy()])
        _, other = self.add([buy(100, 20000)], account="OTHER0001")
        scope = first["scope"]
        report = service.build_report(self.store, scope)
        target = report["imports"][0]
        for target_scope, target_id in ((scope, 999999), (other["scope"], target["id"])):
            with self.assertRaises(service.ReconciliationNotFound):
                service.delete_report(self.store, target_scope, target_id, target["revision"])
        with self.assertRaises(service.ReconciliationNotFound):
            service.delete_scope(self.store, "unknown", report["revision"])
        for revision in (None, "", "a" * 63, "z" * 64):
            with self.subTest(revision=revision), self.assertRaises(ValueError):
                service.delete_report(self.store, scope, target["id"], revision)
            with self.subTest(revision=revision), self.assertRaises(ValueError):
                service.delete_scope(self.store, scope, revision)
        self.assertEqual(service.build_report(self.store, scope), report)
        service.delete_report(self.store, scope, target["id"], target["revision"])
        with self.assertRaises(service.ReconciliationNotFound):
            service.delete_report(self.store, scope, target["id"], target["revision"])

    def test_decimal_source_values_are_preserved_and_no_ledger_or_nav_writes(self):
        self.add([["Mua", "ABC", 3, "0.1", "0.3"], ["Bán", "ABC", 1, "0.2", "0.2"]])
        self.add([["Mua", "ABC", 3, "0.2", "0.6"]], day="06/01/2026")
        position = service.build_report(self.store)["positions"][0]
        self.assertEqual(position["quantity"], 5)
        self.assertAlmostEqual(position["total_cost"], 0.8)
        self.assertAlmostEqual(position["average_cost"], 0.16)
        with self.store.connect() as conn:
            raw = json.loads(conn.execute("SELECT rows_json FROM reconciliation_imports ORDER BY trade_date").fetchone()[0])
            self.assertEqual(raw[0]["buy_amount"], "0.3")
            for table in ("ledger_entries", "ledger_closes", "nav_snapshots"):
                self.assertEqual(conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0], 0)


class FeesAndSellPlanTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = SignalStore(Path(self.temp.name) / "test.db")

    def tearDown(self):
        self.temp.cleanup()

    def add(self, rows, **kwargs):
        parsed = parse(rows, **kwargs)
        service.import_report(self.store, parsed, "tester")
        return parsed

    def test_zero_default_fees_preserve_cost_and_config_does_not_change_source(self):
        parsed = self.add([buy(100, 10000), sell(40)])
        scope = parsed["scope"]
        report = service.build_report(self.store, scope)
        self.assertEqual(report["fees"], dict(scope=scope, buy_fee_percent=0, sell_fee_percent=0, sell_tax_percent=0,
                                             configured=False, updated_by=None, updated_at=None))
        self.assertTrue(report["cost_includes_buy_fee"])
        self.assertEqual(report["positions"][0]["average_cost"], 10000)
        self.assertEqual(report["positions"][0]["total_cost"], 600000)
        self.assertEqual(report["rows"][0]["residual_buy_fee"], 0)
        self.assertEqual(report["rows"][0]["buy_average_price_with_fee"], 10000)
        with self.store.connect() as conn:
            before_source = dict(conn.execute("SELECT * FROM reconciliation_imports").fetchone())
        saved = service.save_fees(self.store, scope, 0, 0, 0, "tester")
        self.assertTrue(saved["configured"])
        configured = service.build_report(self.store, scope)
        self.assertEqual(configured["positions"], report["positions"])
        self.assertEqual(configured["rows"], report["rows"])
        self.assertNotEqual(configured["revision"], report["revision"])
        self.assertEqual(configured["imports"][0]["revision"], report["imports"][0]["revision"])
        with self.store.connect() as conn:
            self.assertEqual(dict(conn.execute("SELECT * FROM reconciliation_imports").fetchone()), before_source)
        self.assertEqual(service.assess_import(self.store, parsed), "duplicate")

    def test_buy_fee_replays_weighted_cost_partial_sale_and_new_purchase_without_touching_base(self):
        parsed = self.add([buy(100, 10000), sell(40)])
        scope = parsed["scope"]
        self.add([buy(100, 20000), sell(50)], day="06/01/2026")
        self.add([sell(80)], day="07/01/2026")
        service.save_base(self.store, scope, "ABC", 10000, "tester")
        service.save_fees(self.store, scope, "0.5", "0.15", "0.1", "tester")
        report = service.build_report(self.store, scope)
        position = report["positions"][0]
        self.assertEqual(position["quantity"], 30)
        self.assertEqual(position["total_quantity"], 10030)
        self.assertAlmostEqual(position["average_cost"], 1608000 / 110)
        self.assertAlmostEqual(position["total_cost"], 30 * 1608000 / 110)
        purchase = next(r for r in report["rows"] if r["trade_date"] == "2026-01-05")
        self.assertEqual(purchase["buy_average_price"], 10000)
        self.assertEqual(purchase["buy_average_price_with_fee"], 10050)
        self.assertEqual(purchase["residual_buy_cost"], 603000)
        self.assertEqual(purchase["residual_buy_fee"], 3000)
        self.assertEqual(report["rows"][0]["closing_average_cost"], report["rows"][0]["opening_average_cost"])
        self.add([buy(100, 25000)], day="08/01/2026")
        position = service.build_report(self.store, scope)["positions"][0]
        self.assertEqual(position["quantity"], 130)
        self.assertAlmostEqual(position["average_cost"], (30 * 1608000 / 110 + 2512500) / 130)
        service.save_fees(self.store, scope, "0.15", "0.2", "0.1", "tester")
        position = service.build_report(self.store, scope)["positions"][0]
        self.assertAlmostEqual(position["average_cost"], (30 * 1600000 / 110 + 2500000) * 1.0015 / 130)
        self.assertEqual(position["base_quantity"], 10000)

    def test_fee_validation_is_atomic_and_rates_are_persisted_as_decimal_text(self):
        parsed = self.add([buy()])
        scope = parsed["scope"]
        service.save_fees(self.store, scope, "0.150000", "0.2", "0.1", "tester")
        before = service.build_report(self.store, scope)
        invalid = (-1, True, None, "NaN", "Infinity", "-Infinity", "100", "0.1234567", "1e-10000", "1e-999999999", "0." + "1" * 130, "text")
        for index in range(3):
            for value in invalid:
                fees = ["0.1", "0.1", "0.1"]
                fees[index] = value
                with self.subTest(index=index, value=value), self.assertRaises(ValueError):
                    service.save_fees(self.store, scope, *fees, "tester")
        with self.assertRaises(ValueError):
            service.save_fees(self.store, scope, 0, 99, 1, "tester")
        with self.assertRaises(ValueError):
            service.save_fees(self.store, "unknown", 0, 0, 0, "tester")
        self.assertEqual(service.build_report(self.store, scope), before)
        with self.store.connect() as conn:
            rates = dict(conn.execute("SELECT * FROM reconciliation_fees").fetchone())
            self.assertEqual(rates["buy_fee_percent"], "0.15")
            self.assertEqual(rates["sell_fee_percent"], "0.2")
            self.assertEqual(rates["sell_tax_percent"], "0.1")
        service.save_fees(self.store, scope, "99.999999", "99.999998", "0.000001", "tester")
        self.assertEqual(service.build_report(self.store, scope)["fees"]["buy_fee_percent"], 99.999999)
        saved = service.save_fees(self.store, scope, "0e-999999999", "0e-999999999", "0e-999999999", "tester")
        self.assertEqual(saved["buy_fee_percent"], 0)
        with self.store.connect() as conn:
            self.assertEqual(conn.execute("SELECT buy_fee_percent FROM reconciliation_fees").fetchone()[0], "0")

    def test_fee_changes_are_isolated_by_scope_and_invalidate_only_account_revision(self):
        first = self.add([buy(100, 10000)])
        other = self.add([buy(100, 10000)], account="OTHER0002")
        other_before = service.build_report(self.store, other["scope"])
        first_before = service.build_report(self.store, first["scope"])
        service.save_fees(self.store, first["scope"], "1", "1", "1", "tester")
        first_after = service.build_report(self.store, first["scope"])
        self.assertNotEqual(first_after["revision"], first_before["revision"])
        self.assertEqual(first_after["imports"][0]["revision"], first_before["imports"][0]["revision"])
        self.assertEqual(first_after["positions"][0]["average_cost"], 10100)
        self.assertEqual(service.build_report(self.store, other["scope"]), other_before)
        with self.assertRaises(service.ReconciliationConflict):
            service.delete_scope(self.store, first["scope"], first_before["revision"])
        self.assertEqual(service.build_report(self.store, first["scope"]), first_after)

    def test_fee_only_scope_survives_single_delete_allows_changes_and_full_clear_removes_fees(self):
        parsed = self.add([buy()])
        scope = parsed["scope"]
        target = service.build_report(self.store, scope)["imports"][0]
        service.save_fees(self.store, scope, "0.15", "0.2", "0.1", "tester")
        service.delete_report(self.store, scope, target["id"], target["revision"])
        report = service.build_report(self.store, scope)
        self.assertEqual(report["imports"], [])
        self.assertEqual(report["positions"], [])
        self.assertEqual(report["scopes"][0]["account_label"], "Tài khoản · #" + scope[-6:])
        self.assertTrue(report["fees"]["configured"])
        service.save_fees(self.store, scope, "0.2", "0.25", "0.1", "tester")
        service.save_base(self.store, scope, "ABC", 10000, "tester")
        report = service.build_report(self.store, scope)
        result = service.delete_scope(self.store, scope, report["revision"])
        self.assertEqual(result, dict(scope=scope, deleted_imports_count=0, deleted_notes_count=0, deleted_bases_count=1, deleted_fees_count=1))
        self.assertIsNone(service.build_report(self.store)["selected_scope"])
        with self.store.connect() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM reconciliation_fees").fetchone()[0], 0)

    def test_sell_plan_uses_total_surplus_equal_parts_after_fees_and_preserves_all_data(self):
        parsed = self.add([buy(10, 10000)])
        scope = parsed["scope"]
        service.save_base(self.store, scope, "ABC", 10000, "tester")
        service.save_fees(self.store, scope, 1, 1, 1, "tester")
        before = service.build_report(self.store, scope)
        plan = service.build_sell_plan(self.store, scope, "abc", 3, 10, 20, 100)
        self.assertEqual(plan["quantity"], 10)
        self.assertEqual(plan["average_cost"], 10100)
        self.assertEqual(plan["total_cost"], 101000)
        self.assertEqual([o["quantity"] for o in plan["orders"]], [4, 3, 3])
        self.assertEqual([o["target_profit_percent"] for o in plan["orders"]], [10, 15, 20])
        self.assertEqual([o["price"] for o in plan["orders"]], [11400, 11900, 12400])
        self.assertEqual(plan["price_min"], 11400)
        self.assertEqual(plan["price_max"], 12400)
        first = plan["orders"][0]
        self.assertEqual(first["allocated_cost"], 40400)
        self.assertEqual(first["gross_proceeds"], 45600)
        self.assertEqual(first["sell_fee"], 456)
        self.assertEqual(first["sell_tax"], 456)
        self.assertEqual(first["net_proceeds"], 44688)
        self.assertEqual(first["net_profit"], 4288)
        self.assertAlmostEqual(first["net_profit_percent"], 4288 / 40400 * 100)
        totals = plan["totals"]
        self.assertEqual(totals["quantity"], 10)
        self.assertEqual(totals["allocated_cost"], 101000)
        self.assertEqual(totals["gross_proceeds"], 118500)
        self.assertEqual(totals["sell_fee"], 1185)
        self.assertEqual(totals["sell_tax"], 1185)
        self.assertEqual(totals["net_proceeds"], 116130)
        self.assertEqual(totals["net_profit"], 15130)
        self.assertAlmostEqual(totals["net_profit_percent"], 15130 / 101000 * 100)
        self.assertEqual(plan["revision"], before["revision"])
        for order in plan["orders"]:
            self.assertGreaterEqual(order["net_profit_percent"], order["target_profit_percent"])
        self.assertEqual(service.build_report(self.store, scope), before)

    def test_multiple_buy_days_split_current_total_not_daily_cohorts(self):
        parsed = self.add([buy(3, 10000)])
        self.add([buy(4, 20000)], day="06/01/2026")
        plan = service.build_sell_plan(self.store, parsed["scope"], "ABC", 3, 1, 3)
        self.assertEqual(plan["quantity"], 7)
        self.assertEqual([o["quantity"] for o in plan["orders"]], [3, 2, 2])
        self.assertEqual([o["price"] for o in plan["orders"]], [15900, 16100, 16200])
        self.assertEqual(plan["totals"]["allocated_cost"], 110000)
        self.assertAlmostEqual(plan["average_cost"], 110000 / 7)

    def test_one_part_uses_minimum_target_but_band_keeps_requested_upper_target(self):
        parsed = self.add([buy(10, 10000)])
        service.save_fees(self.store, parsed["scope"], 1, 1, 1, "tester")
        plan = service.build_sell_plan(self.store, parsed["scope"], "ABC", 1, 10, 20)
        self.assertEqual(plan["orders"][0]["quantity"], 10)
        self.assertEqual(plan["orders"][0]["price"], 11400)
        self.assertEqual(plan["orders"][0]["target_profit_percent"], 10)
        self.assertEqual(plan["price_min"], 11400)
        self.assertEqual(plan["price_max"], 12400)
        fixed = service.build_sell_plan(self.store, parsed["scope"], "ABC", 2, "10")
        self.assertEqual(fixed["price_min"], fixed["price_max"])
        self.assertEqual([o["price"] for o in fixed["orders"]], [11400, 11400])

    def test_exact_ceiling_handles_repeating_average_partial_sale_and_fee_boundaries(self):
        parsed = self.add([["Mua", "ABC", 3, "0.666666666667", "2"]])
        scope = parsed["scope"]
        plan = service.build_sell_plan(self.store, scope, "ABC", 1, 50, price_step=1)
        self.assertEqual(plan["price_min"], 1)
        self.add([sell(1, 1)], day="06/01/2026")
        plan = service.build_sell_plan(self.store, scope, "ABC", 2, 50, price_step=1)
        self.assertEqual([o["price"] for o in plan["orders"]], [1, 1])
        self.assertEqual(plan["totals"]["net_profit_percent"], 50)
        self.add([buy(1, 1)], day="07/01/2026")
        plan = service.build_sell_plan(self.store, scope, "ABC", 1, 800, price_step=1)
        self.assertEqual(plan["price_min"], 7)
        fee_boundary = self.add([buy(1, 99.9)], account="OTHER0002")
        service.save_fees(self.store, fee_boundary["scope"], 0, 0, "0.1", "tester")
        self.assertEqual(service.build_sell_plan(self.store, fee_boundary["scope"], "ABC", 1, 0, price_step=100)["price_min"], 100)

    def test_ceiling_does_not_lose_source_precision_via_public_float(self):
        above = self.add([["Mua", "ABC", 1, "100.000000000000000000001", "100.000000000000000000001"]])
        below = self.add([["Mua", "ABC", 1, "99.999999999999999999999", "99.999999999999999999999"]], account="OTHER0002")
        self.assertEqual(service.build_report(self.store, above["scope"])["positions"][0]["average_cost"], 100)
        self.assertEqual(service.build_sell_plan(self.store, above["scope"], "ABC", 1, 0)["price_min"], 200)
        self.assertEqual(service.build_sell_plan(self.store, below["scope"], "ABC", 1, 0)["price_min"], 100)

    def test_planner_validation_rejects_base_only_missing_holdings_and_invalid_parameters(self):
        parsed = self.add([buy(3, 10000)])
        scope = parsed["scope"]
        service.save_base(self.store, scope, "XYZ", 10000, "tester")
        before = service.build_report(self.store, scope)
        invalid = [dict(parts=p) for p in (0, 4, 101, True, 1.0, "1")]
        invalid += [dict(price_step=s) for s in (0, 1000001, True, 1.5, "100")]
        invalid += [dict(target_profit_percent=p) for p in (-1, True, None, "NaN", "Infinity", 10001, "0.1234567", "1e-999999999", "0." + "1" * 130)]
        invalid += [dict(upper_profit_percent=p) for p in (-1, True, "NaN", 10001, "0.1234567")]
        invalid += [dict(ticker=t) for t in ("XYZ", "MISSING", "<script>", None)]
        invalid += [dict(scope="unknown")]
        for changed in invalid:
            kwargs = dict(scope=scope, ticker="ABC", parts=1, target_profit_percent=1, upper_profit_percent=2, price_step=100)
            kwargs.update(changed)
            with self.subTest(changed=changed), self.assertRaises(ValueError):
                service.build_sell_plan(self.store, **kwargs)
        with self.assertRaises(ValueError):
            service.build_sell_plan(self.store, scope, "ABC", 1, 10, 9)
        self.assertEqual(service.build_report(self.store, scope), before)


if __name__ == "__main__":
    unittest.main()
