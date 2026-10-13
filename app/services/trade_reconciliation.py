"""Trading inventory above a separately held base, independent of the ledger.

Daily buy-minus-sell quantities add to or consume prior trading inventory.
Sales keep its weighted average cost and cannot reduce the fixed, already
owned base. Values and average costs are in VND. Configured buy fees apply
to surplus cost; sell fees and tax apply to suggested prices, not quantities.
"""
from __future__ import annotations

import hashlib
import json
import re
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, localcontext
from fractions import Fraction
from io import BytesIO
from pathlib import PurePath

from openpyxl import load_workbook

from app.database import utc_now_iso
from app.services.market_hours import MARKET_TZ


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ROWS = 20000
MAX_COLUMNS = 50
MAX_QUANTITY = 10**12
MAX_AMOUNT = Decimal("1e18")
HEADERS = {
    "giao dich": "side", "ma": "ticker", "khoi luong khop": "quantity",
    "gia khop binh quan": "average", "gia tri khop": "amount",
}


class ReconciliationConflict(ValueError):
    """An import conflicts or a deletion snapshot has changed."""


class ReconciliationNotFound(ValueError):
    """The selected imported report or account no longer exists."""


def _normalize(value) -> str:
    text = str(value or "").replace("\xa0", " ").strip().casefold()
    text = unicodedata.normalize("NFD", text).replace("đ", "d")
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    return " ".join(text.split())


def _decimal(value, label: str, *, allow_zero=False) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"{label}: cần số hợp lệ")
    text = str(value).strip().replace("\xa0", "").replace(" ", "")
    if len(text) > 128:
        raise ValueError(f"{label}: định dạng số vượt giới hạn")
    # The supplied broker export uses commas for thousands and a decimal dot.
    if "," in text:
        if not re.fullmatch(r"[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?", text):
            raise ValueError(f"{label}: định dạng số không hợp lệ")
        text = text.replace(",", "")
    try:
        number = Decimal(text)
    except InvalidOperation:
        raise ValueError(f"{label}: cần số hợp lệ") from None
    if (not number.is_finite() or number < 0 or (not allow_zero and number == 0)
            or number > MAX_AMOUNT or (number and number < Decimal("1e-12"))):
        raise ValueError(f"{label}: số phải hữu hạn, {'không âm' if allow_zero else 'lớn hơn 0'} và trong giới hạn")
    return number


def _exact(value: Decimal) -> str:
    return format(value.normalize(), "f") if value else "0"


def _safe_filename(filename: str) -> str:
    name = PurePath(str(filename).replace("\\", "/")).name
    name = "".join(c for c in name if c.isprintable())[:200]
    if not name.casefold().endswith(".xlsx"):
        raise ValueError("Chỉ hỗ trợ báo cáo Excel .xlsx")
    return name


def _validate_archive(data: bytes):
    if not isinstance(data, bytes) or not data or len(data) > MAX_FILE_BYTES:
        raise ValueError("Báo cáo phải có dữ liệu và không vượt quá 10 MB")
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            items = archive.infolist()
            if len(items) > 1000 or sum(i.file_size for i in items) > MAX_EXPANDED_BYTES:
                raise ValueError("Nội dung Excel vượt giới hạn giải nén 50 MB")
            if any(i.flag_bits & 1 or i.file_size > 20 * 1024 * 1024 or
                   (i.file_size > 1024 * 1024 and i.file_size / max(i.compress_size, 1) > 1000)
                   for i in items):
                raise ValueError("Cấu trúc Excel không được hỗ trợ hoặc vượt giới hạn")
            names = {i.filename.casefold() for i in items}
            if "xl/workbook.xml" not in names or any("vbaproject" in n or n.startswith("xl/externallinks/") for n in names):
                raise ValueError("Chỉ hỗ trợ workbook .xlsx không có macro/liên kết ngoài")
    except (zipfile.BadZipFile, OSError):
        raise ValueError("Tệp không phải workbook Excel .xlsx hợp lệ") from None


def _read_metadata(sheet, header_row):
    metadata = {}
    totals = {}
    labels = {
        "ngay giao dich": "date", "so tai khoan": "account", "tai khoan": "account",
        "tieu khoan": "subaccount", "tong gia tri mua": "buy_amount",
        "tong gia tri ban": "sell_amount", "tong gia tri khop": "all_amount",
    }
    for row in sheet.iter_rows(min_row=1, max_row=header_row - 1):
        for cell in row:
            if cell.value is None:
                continue
            if cell.data_type == "f":
                raise ValueError("Thông tin báo cáo không được dùng công thức Excel")
            text = str(cell.value).strip()
            if ":" not in text:
                continue
            label, value = text.split(":", 1)
            key = labels.get(_normalize(label))
            if not key:
                continue
            # Some exports place metadata in the adjacent cell instead.
            if not value.strip() and cell.column < sheet.max_column:
                adjacent = sheet.cell(cell.row, cell.column + 1)
                if adjacent.data_type == "f":
                    raise ValueError("Thông tin báo cáo không được dùng công thức Excel")
                value = str(adjacent.value or "")
            target = totals if key.endswith("amount") else metadata
            if key in target and str(target[key]).strip() != value.strip():
                raise ValueError("Thông tin đầu báo cáo không nhất quán")
            target[key] = value.strip()
    if not metadata.get("date"):
        raise ValueError("Thiếu Ngày giao dịch trong báo cáo; không suy đoán ngày từ thời điểm tải lên")
    raw_date = metadata["date"].split()[0]
    try:
        trade_date = datetime.strptime(raw_date, "%d/%m/%Y").date()
    except ValueError:
        try:
            trade_date = date.fromisoformat(raw_date)
        except ValueError:
            raise ValueError("Ngày giao dịch không hợp lệ; cần dd/mm/yyyy") from None
    if trade_date > datetime.now(MARKET_TZ).date():
        raise ValueError("Ngày giao dịch không được nằm trong tương lai")
    account = re.sub(r"\s+", "", metadata.get("account", ""))
    if not re.fullmatch(r"[A-Za-z0-9._/-]{3,64}", account):
        raise ValueError("Thiếu hoặc sai Số tài khoản trong báo cáo")
    subaccount = metadata.get("subaccount", "Không ghi tiểu khoản")
    if not subaccount or len(subaccount) > 80:
        raise ValueError("Tiểu khoản không hợp lệ")
    identity = account.casefold() + "\n" + _normalize(subaccount)
    scope = "acct_" + hashlib.sha256(identity.encode()).hexdigest()[:24]
    label = "••••" + account[-4:] + " · " + subaccount
    return trade_date.isoformat(), scope, label, totals


def _canonical_rows(rows):
    output = []
    seen = set()
    for row in rows:
        ticker = str(row["ticker"]).upper()
        if ticker in seen or not re.fullmatch(r"[A-Z0-9][A-Z0-9._-]{0,19}", ticker):
            raise ValueError("Mã cổ phiếu không hợp lệ hoặc bị trùng trong dữ liệu tổng hợp")
        seen.add(ticker)
        quantities = []
        amounts = []
        for side in ("buy", "sell"):
            q = _decimal(row[f"{side}_quantity"], "Khối lượng", allow_zero=True)
            if q != q.to_integral_value() or q > MAX_QUANTITY:
                raise ValueError("Khối lượng khớp phải là số nguyên trong giới hạn")
            amount = _decimal(row.get(f"{side}_amount_exact", row[f"{side}_amount"]), "Giá trị khớp", allow_zero=True)
            if bool(q) != bool(amount):
                raise ValueError("Khối lượng và giá trị khớp không nhất quán")
            quantities.append(int(q)); amounts.append(_exact(amount))
        if not any(quantities):
            raise ValueError("Dòng mã phải có ít nhất một giao dịch")
        output.append(dict(ticker=ticker, buy_quantity=quantities[0], sell_quantity=quantities[1],
                           buy_amount=amounts[0], sell_amount=amounts[1]))
    if not output or len(output) > MAX_ROWS:
        raise ValueError("Báo cáo không có giao dịch hoặc có quá nhiều mã")
    return sorted(output, key=lambda r: r["ticker"])


def _digest(scope, trade_date, canonical):
    payload = dict(scope=scope, trade_date=trade_date, rows=canonical)
    return hashlib.sha256(json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def _revision(payload):
    """Hash persisted content and audit fields, never computed holdings."""
    return hashlib.sha256(json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def _report_revision(item, notes):
    report_notes = [n for n in notes if n["scope"] == item["scope"] and n["trade_date"] == item["trade_date"]]
    return _revision(dict(report=item, notes=sorted(report_notes, key=lambda n: n["ticker"])))


def _scope_revision(scope, imports, notes, bases, fees):
    return _revision(dict(scope=scope,
                          imports=sorted((i for i in imports if i["scope"] == scope), key=lambda i: (i["trade_date"], i["id"])),
                          notes=sorted((n for n in notes if n["scope"] == scope), key=lambda n: (n["trade_date"], n["ticker"])),
                          bases=sorted((b for b in bases if b["scope"] == scope), key=lambda b: b["ticker"]),
                          fees=[f for f in fees if f["scope"] == scope]))


def _scope_data(conn, scope):
    return tuple([dict(r) for r in conn.execute(f"SELECT * FROM {table} WHERE scope=?", (scope,))]
                 for table in ("reconciliation_imports", "reconciliation_notes", "reconciliation_bases", "reconciliation_fees"))


def _scope_exists(conn, scope):
    return conn.execute("SELECT 1 FROM reconciliation_imports WHERE scope=? UNION ALL "
                        "SELECT 1 FROM reconciliation_bases WHERE scope=? UNION ALL "
                        "SELECT 1 FROM reconciliation_fees WHERE scope=? LIMIT 1", (scope, scope, scope)).fetchone() is not None


def _percent(value, label, *, maximum=Decimal(100), inclusive=False):
    if isinstance(value, bool) or value is None:
        raise ValueError(f"{label}: cần tỷ lệ phần trăm hợp lệ")
    text = str(value).strip()
    if len(text) > 128:
        raise ValueError(f"{label}: định dạng số vượt giới hạn")
    try:
        number = Decimal(text)
    except InvalidOperation:
        raise ValueError(f"{label}: cần tỷ lệ phần trăm hợp lệ") from None
    if not number.is_finite() or number < 0 or number > maximum or (not inclusive and number == maximum):
        bound = "không vượt quá" if inclusive else "nhỏ hơn"
        raise ValueError(f"{label}: tỷ lệ phải không âm và {bound} {maximum}%")
    if not number:
        return Decimal(0)
    # Inspect the tuple without normalize(): Decimal context can underflow
    # extreme negative exponents to zero before validation.
    digits = number.as_tuple().digits
    trailing_zeroes = 0
    for digit in reversed(digits):
        if digit:
            break
        trailing_zeroes += 1
    if number.as_tuple().exponent + trailing_zeroes < -6:
        raise ValueError(f"{label}: tối đa 6 chữ số thập phân")
    return number


def _public_fees(scope, record=None):
    return dict(scope=scope, **{key: float(Decimal(record[key])) if record else 0.0
                               for key in ("buy_fee_percent", "sell_fee_percent", "sell_tax_percent")},
                configured=record is not None, updated_by=record["updated_by"] if record else None,
                updated_at=record["updated_at"] if record else None)


def _check_revision(expected, current):
    if not isinstance(expected, str) or not re.fullmatch(r"[a-f0-9]{64}", expected):
        raise ValueError("Cần phiên bản dữ liệu hợp lệ trước khi xóa; hãy tải lại dữ liệu")
    if expected != current:
        raise ReconciliationConflict("Dữ liệu đã thay đổi; hãy tải lại và kiểm tra trước khi xóa")


def _public_row(row):
    buy, sell = row["buy_quantity"], row["sell_quantity"]
    ba, sa = Decimal(row["buy_amount"]), Decimal(row["sell_amount"])
    return dict(ticker=row["ticker"], buy_quantity=buy, sell_quantity=sell,
                buy_amount=float(ba), sell_amount=float(sa),
                buy_amount_exact=_exact(ba), sell_amount_exact=_exact(sa),
                buy_average_price=float(ba / buy) if buy else None,
                sell_average_price=float(sa / sell) if sell else None,
                matched_quantity=min(buy, sell), net_quantity=buy - sell,
                residual_buy_quantity=max(buy - sell, 0))


def parse_workbook(data: bytes, filename: str) -> dict:
    """Read a complete daily broker export; never retain its account number/blob."""
    filename = _safe_filename(filename)
    _validate_archive(data)
    try:
        # Normal mode ignores the supplied export's incorrect dimension='A1'.
        workbook = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=False)
    except Exception:
        raise ValueError("Không đọc được workbook Excel; hãy xuất lại đúng mẫu báo cáo khớp lệnh") from None
    try:
        candidates = []
        if len(workbook.worksheets) > 50:
            raise ValueError("Workbook có quá nhiều sheet")
        for sheet in workbook.worksheets:
            if sheet.max_row > MAX_ROWS + 100 or sheet.max_column > MAX_COLUMNS:
                raise ValueError("Sheet vượt giới hạn số dòng/cột của mẫu báo cáo")
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 100)):
                mapping = {}
                for cell in row:
                    key = HEADERS.get(_normalize(cell.value))
                    if key:
                        if key in mapping:
                            raise ValueError("Cột tiêu đề bị trùng trong báo cáo")
                        mapping[key] = cell.column
                if len(mapping) == len(HEADERS):
                    candidates.append((sheet, row[0].row, mapping))
        if len(candidates) != 1:
            raise ValueError("Cần đúng một bảng với cột Giao dịch, Mã, Khối lượng khớp, Giá khớp bình quân, Giá trị khớp")
        sheet, header_row, columns = candidates[0]
        trade_date, scope, label, declared_totals = _read_metadata(sheet, header_row)
        grouped = {}
        for row_number in range(header_row + 1, sheet.max_row + 1):
            cells = {k: sheet.cell(row_number, c) for k, c in columns.items()}
            if all(c.value is None or str(c.value).strip() == "" for c in cells.values()):
                continue
            if any(c.data_type == "f" for c in cells.values()):
                raise ValueError(f"Dòng {row_number}: dữ liệu khớp lệnh không được dùng công thức Excel")
            side = {"mua": "buy", "ban": "sell"}.get(_normalize(cells["side"].value))
            if not side:
                raise ValueError(f"Dòng {row_number}: Giao dịch phải là Mua hoặc Bán")
            ticker = str(cells["ticker"].value or "").strip().upper()
            if not re.fullmatch(r"[A-Z0-9][A-Z0-9._-]{0,19}", ticker):
                raise ValueError(f"Dòng {row_number}: mã cổ phiếu không hợp lệ")
            quantity = _decimal(cells["quantity"].value, f"Dòng {row_number}, khối lượng")
            if quantity != quantity.to_integral_value() or quantity > MAX_QUANTITY:
                raise ValueError(f"Dòng {row_number}: khối lượng khớp phải là số nguyên trong giới hạn")
            amount = _decimal(cells["amount"].value, f"Dòng {row_number}, giá trị khớp")
            average = _decimal(cells["average"].value, f"Dòng {row_number}, giá bình quân")
            # The export rounds averages to six decimals. A VND 0.01 tolerance
            # admits that display rounding without hiding meaningful mismatch.
            with localcontext() as context:
                context.prec = 50
                if abs(average - amount / quantity) > Decimal("0.01"):
                    raise ValueError(f"Dòng {row_number}: giá trị khớp không khớp khối lượng × giá bình quân")
            aggregate = grouped.setdefault(ticker, dict(ticker=ticker, buy_quantity=0, sell_quantity=0,
                                                        buy_amount=Decimal(0), sell_amount=Decimal(0)))
            aggregate[f"{side}_quantity"] += int(quantity)
            aggregate[f"{side}_amount"] += amount
        canonical = _canonical_rows(grouped.values())
        totals_exact = {side: sum((Decimal(r[f"{side}_amount"]) for r in canonical), Decimal(0))
                        for side in ("buy", "sell")}
        for key, declared in declared_totals.items():
            expected = sum(totals_exact.values()) if key == "all_amount" else totals_exact[key.split("_")[0]]
            if _decimal(declared, "Tổng giá trị báo cáo", allow_zero=True) != expected:
                raise ValueError("Tổng giá trị đầu báo cáo không khớp các dòng khớp lệnh")
        return dict(trade_date=trade_date, scope=scope, account_label=label, filename=filename,
                    digest=_digest(scope, trade_date, canonical), rows=[_public_row(r) for r in canonical],
                    totals={**{f"{s}_quantity": sum(r[f"{s}_quantity"] for r in canonical) for s in ("buy", "sell")},
                            **{f"{s}_amount": float(v) for s, v in totals_exact.items()}}, currency="VND")
    finally:
        workbook.close()


def assess_import(store, parsed: dict) -> str:
    canonical = _canonical_rows(parsed["rows"])
    digest = _digest(parsed["scope"], parsed["trade_date"], canonical)
    with store.connect() as conn:
        row = conn.execute("SELECT digest FROM reconciliation_imports WHERE scope=? AND trade_date=?",
                           (parsed["scope"], parsed["trade_date"])).fetchone()
    return "new" if row is None else "duplicate" if row["digest"] == digest else "conflict"


def import_report(store, parsed: dict, username, replace_existing=False) -> dict:
    canonical = _canonical_rows(parsed["rows"])
    scope, day = parsed["scope"], parsed["trade_date"]
    digest = _digest(scope, day, canonical)
    encoded = json.dumps(canonical, sort_keys=True, separators=(",", ":"))
    with store.connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        existing = conn.execute("SELECT id,digest FROM reconciliation_imports WHERE scope=? AND trade_date=?", (scope, day)).fetchone()
        duplicate = existing is not None and existing["digest"] == digest
        replaced = existing is not None and not duplicate
        if replaced and not replace_existing:
            raise ReconciliationConflict("Tài khoản/ngày đã có báo cáo khác; xem trước và chọn thay thế toàn bộ báo cáo ngày đó")
        if not duplicate:
            values = (parsed["account_label"], _safe_filename(parsed["filename"]), digest, encoded, username, utc_now_iso())
            if replaced:
                conn.execute("UPDATE reconciliation_imports SET account_label=?,filename=?,digest=?,rows_json=?,created_by=?,created_at=? WHERE id=?",
                             (*values, existing["id"]))
                tickers = [r["ticker"] for r in canonical]
                if tickers:
                    placeholders = ",".join("?" for _ in tickers)
                    conn.execute(f"DELETE FROM reconciliation_notes WHERE scope=? AND trade_date=? AND ticker NOT IN ({placeholders})", (scope, day, *tickers))
                else:
                    conn.execute("DELETE FROM reconciliation_notes WHERE scope=? AND trade_date=?", (scope, day))
            else:
                conn.execute("INSERT INTO reconciliation_imports(scope,trade_date,account_label,filename,digest,rows_json,created_by,created_at) VALUES(?,?,?,?,?,?,?,?)",
                             (scope, day, *values))
    return dict(inserted=int(not duplicate), replaced=replaced, duplicate=duplicate,
                scope=scope, trade_date=day, rows_count=len(canonical))


def _build_report_snapshot(store, scope=None):
    with store.connect() as conn:
        # A deferred read transaction keeps the exposed revisions consistent
        # with the report, even if another session edits notes or imports.
        conn.execute("BEGIN")
        imports = [dict(r) for r in conn.execute("SELECT * FROM reconciliation_imports ORDER BY trade_date,id")]
        all_notes = [dict(r) for r in conn.execute("SELECT * FROM reconciliation_notes")]
        notes = {(r["scope"], r["trade_date"], r["ticker"]): r["note"] for r in all_notes}
        all_bases = [dict(r) for r in conn.execute("SELECT * FROM reconciliation_bases ORDER BY scope,ticker")]
        all_fees = [dict(r) for r in conn.execute("SELECT * FROM reconciliation_fees ORDER BY scope")]
    scope_map = {}
    for item in imports:
        s = scope_map.setdefault(item["scope"], dict(scope=item["scope"], account_label=item["account_label"],
                                                   first_trade_date=item["trade_date"], last_trade_date=item["trade_date"], imports_count=0))
        s["last_trade_date"] = item["trade_date"]
        s["imports_count"] += 1
    for record in [*all_bases, *all_fees]:
        scope_map.setdefault(record["scope"], dict(scope=record["scope"], account_label="Tài khoản · #" + record["scope"][-6:],
                                                first_trade_date=None, last_trade_date=None, imports_count=0))
    if scope and scope not in scope_map:
        raise ValueError("Không tìm thấy tài khoản đối chiếu")
    scopes = sorted(scope_map.values(), key=lambda s: (s["account_label"], s["scope"]))
    label_counts = {}
    for item in scopes:
        label_counts[item["account_label"]] = label_counts.get(item["account_label"], 0) + 1
    for item in scopes:
        if label_counts[item["account_label"]] > 1:
            # Masked account numbers may collide; expose a stable hash suffix
            # only in the selector label, without retaining the full account.
            item["account_label"] += " · #" + item["scope"][-6:]
    selected_scope = scope or (scopes[0]["scope"] if scopes else None)
    selected = [i for i in imports if i["scope"] == selected_scope]
    bases = [b for b in all_bases if b["scope"] == selected_scope]
    fee_record = next((f for f in all_fees if f["scope"] == selected_scope), None)
    fees = _public_fees(selected_scope, fee_record)
    buy_fee = Fraction(Decimal(fee_record["buy_fee_percent"])) / 100 if fee_record else Fraction(0)
    base_quantities = {b["ticker"]: b["base_quantity"] for b in bases}
    holdings = {}
    selected_label = scope_map[selected_scope]["account_label"] if selected_scope else ""
    for base in bases:
        key = (selected_scope, base["ticker"])
        holdings[key] = dict(scope=selected_scope, account_label=selected_label, ticker=base["ticker"],
                             quantity=0, cost=Fraction(0), first_trade_date=None, last_trade_date=None)
    rows = []
    ignored = 0
    ignored_sales = 0
    with localcontext() as context:
        context.prec = 50
        for item in selected:
            for raw in json.loads(item["rows_json"]):
                row = _public_row(raw)
                surplus = row["residual_buy_quantity"]
                key = (item["scope"], raw["ticker"])
                position = holdings.setdefault(key, dict(scope=key[0], account_label=item["account_label"], ticker=key[1],
                                                         quantity=0, cost=Fraction(0), first_trade_date=item["trade_date"],
                                                         last_trade_date=item["trade_date"]))
                opening = position["quantity"]
                opening_average = position["cost"] / opening if opening else None
                sold = min(max(-row["net_quantity"], 0), opening)
                ignored_sell = max(-row["net_quantity"] - sold, 0)
                ignored_sales += ignored_sell
                if not surplus and not sold:
                    ignored += 1
                added_cost = Fraction(0)
                added_fee = Fraction(0)
                if surplus:
                    before_fee = Fraction(Decimal(raw["buy_amount"])) * surplus / raw["buy_quantity"]
                    added_fee = before_fee * buy_fee
                    added_cost = before_fee + added_fee
                    position["quantity"] += surplus
                    position["cost"] += added_cost
                elif sold:
                    position["quantity"] -= sold
                    position["cost"] = opening_average * position["quantity"] if position["quantity"] else Fraction(0)
                if position["first_trade_date"] is None:
                    position["first_trade_date"] = item["trade_date"]
                position["last_trade_date"] = item["trade_date"]
                closing_average = position["cost"] / position["quantity"] if position["quantity"] else None
                base_quantity = base_quantities.get(key[1], 0)
                status = "buy_remaining" if position["quantity"] else "base_only" if base_quantity else "closed"
                rows.append({**row, "scope": key[0], "account_label": item["account_label"], "trade_date": item["trade_date"],
                             "opening_quantity": opening, "opening_average_cost": float(opening_average) if opening_average is not None else None,
                             "closing_quantity": position["quantity"], "closing_average_cost": float(closing_average) if closing_average is not None else None,
                             "residual_buy_cost": float(added_cost), "tracked_sell_quantity": sold,
                             "residual_buy_fee": float(added_fee),
                             "buy_average_price_with_fee": float(Fraction(Decimal(raw["buy_amount"])) / raw["buy_quantity"] * (1 + buy_fee)) if raw["buy_quantity"] else None,
                             "buy_fee_percent": fees["buy_fee_percent"],
                             "ignored_sell_quantity": ignored_sell, "base_quantity": base_quantity,
                             "opening_total_quantity": base_quantity + opening,
                             "closing_total_quantity": base_quantity + position["quantity"], "status": status,
                             "note": notes.get((key[0], item["trade_date"], key[1]), "")})
        positions = [{k: v for k, v in p.items() if k != "cost"} |
                     dict(average_cost=float(p["cost"] / p["quantity"]) if p["quantity"] else None,
                          total_cost=float(p["cost"]), base_quantity=base_quantities.get(p["ticker"], 0),
                          total_quantity=p["quantity"] + base_quantities.get(p["ticker"], 0),
                          status="buy_remaining" if p["quantity"] else "base_only" if base_quantities.get(p["ticker"], 0) else "closed")
                     for p in holdings.values()]
        total_cost = float(sum((p["cost"] for p in holdings.values()), Fraction(0)))
    metadata = [{k: v for k, v in i.items() if k not in {"rows_json", "digest"}} |
                dict(rows_count=len(json.loads(i["rows_json"])), revision=_report_revision(i, all_notes)) for i in selected]
    return dict(scopes=scopes, selected_scope=selected_scope, bases=bases, fees=fees,
                revision=_scope_revision(selected_scope, imports, all_notes, all_bases, all_fees) if selected_scope else None,
                imports=list(reversed(metadata)), rows=sorted(rows, key=lambda r: (r["trade_date"], r["ticker"], r["scope"]), reverse=True),
                positions=sorted(positions, key=lambda p: (p["ticker"], p["scope"])),
                totals=dict(imports_count=len(selected), daily_rows_count=len(rows), ignored_rows_count=ignored,
                            outstanding_tickers=len({p["ticker"] for p in positions if p["quantity"] > 0}),
                            outstanding_quantity=sum(p["quantity"] for p in positions), total_cost=total_cost,
                            ignored_sell_quantity=ignored_sales,
                            base_quantity=sum(p["base_quantity"] for p in positions),
                            total_quantity=sum(p["total_quantity"] for p in positions)),
                currency="VND", method="daily_net_inventory_with_base", cost_includes_buy_fee=True), holdings, fee_record


def build_report(store, scope=None) -> dict:
    report, _, _ = _build_report_snapshot(store, scope)
    return report


def save_note(store, scope, trade_date, ticker, note, username) -> dict:
    ticker = str(ticker).strip().upper()
    if not isinstance(note, str) or len(note) > 2000:
        raise ValueError("Ghi chú phải là văn bản tối đa 2000 ký tự")
    note = note.strip()
    updated = utc_now_iso()
    with store.connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        report = conn.execute("SELECT rows_json FROM reconciliation_imports WHERE scope=? AND trade_date=?", (scope, trade_date)).fetchone()
        target = next((r for r in json.loads(report["rows_json"]) if r["ticker"] == ticker), None) if report else None
        if not target:
            raise ValueError("Chỉ ghi chú mã có giao dịch trong báo cáo ngày đã nhập")
        conn.execute("INSERT INTO reconciliation_notes(scope,trade_date,ticker,note,updated_by,updated_at) VALUES(?,?,?,?,?,?) "
                     "ON CONFLICT(scope,trade_date,ticker) DO UPDATE SET note=excluded.note,updated_by=excluded.updated_by,updated_at=excluded.updated_at",
                     (scope, trade_date, ticker, note, username, updated))
    return dict(scope=scope, trade_date=trade_date, ticker=ticker, note=note, updated_by=username, updated_at=updated)


def save_base(store, scope, ticker, quantity, username) -> dict:
    """Set already owned shares held separately from all imported trading fills."""
    if not isinstance(scope, str) or not scope:
        raise ValueError("Cần chọn tài khoản đối chiếu đã có báo cáo")
    if not isinstance(ticker, str):
        raise ValueError("Mã cổ phiếu cơ sở không hợp lệ")
    ticker = ticker.strip().upper()
    if not re.fullmatch(r"[A-Z0-9][A-Z0-9._-]{0,19}", ticker):
        raise ValueError("Mã cổ phiếu cơ sở không hợp lệ")
    if isinstance(quantity, bool) or not isinstance(quantity, int) or not 0 <= quantity <= MAX_QUANTITY:
        raise ValueError("Số cổ phiếu cơ sở phải là số nguyên không âm và không vượt quá 1.000.000.000.000 CP")
    updated = utc_now_iso()
    with store.connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        if not _scope_exists(conn, scope):
            raise ValueError("Không tìm thấy tài khoản đối chiếu đã có dữ liệu")
        conn.execute("INSERT INTO reconciliation_bases(scope,ticker,base_quantity,updated_by,updated_at) VALUES(?,?,?,?,?) "
                     "ON CONFLICT(scope,ticker) DO UPDATE SET base_quantity=excluded.base_quantity,updated_by=excluded.updated_by,updated_at=excluded.updated_at",
                     (scope, ticker, quantity, username, updated))
    return dict(scope=scope, ticker=ticker, base_quantity=quantity, updated_by=username, updated_at=updated)


def delete_report(store, scope, import_id, expected_revision) -> dict:
    """Delete exactly one imported day and its notes, preserving every base."""
    if not isinstance(scope, str) or not scope:
        raise ValueError("Cần chọn tài khoản đối chiếu")
    if isinstance(import_id, bool) or not isinstance(import_id, int) or import_id <= 0:
        raise ValueError("Báo cáo cần xóa không hợp lệ")
    with store.connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute("SELECT * FROM reconciliation_imports WHERE scope=? AND id=?", (scope, import_id)).fetchone()
        if row is None:
            raise ReconciliationNotFound("Không tìm thấy báo cáo cần xóa trong tài khoản đã chọn")
        item = dict(row)
        notes = [dict(r) for r in conn.execute("SELECT * FROM reconciliation_notes WHERE scope=? AND trade_date=?",
                                               (scope, item["trade_date"]))]
        _check_revision(expected_revision, _report_revision(item, notes))
        deleted_notes = conn.execute("DELETE FROM reconciliation_notes WHERE scope=? AND trade_date=?",
                                     (scope, item["trade_date"])).rowcount
        deleted_imports = conn.execute("DELETE FROM reconciliation_imports WHERE scope=? AND id=?", (scope, import_id)).rowcount
    return dict(scope=scope, import_id=import_id, trade_date=item["trade_date"],
                deleted_imports_count=deleted_imports, deleted_notes_count=deleted_notes, deleted_bases_count=0)


def delete_scope(store, scope, expected_revision) -> dict:
    """Clear one account's imports, notes, bases and fee configuration."""
    if not isinstance(scope, str) or not scope:
        raise ValueError("Cần chọn tài khoản đối chiếu")
    with store.connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        imports, notes, bases, fees = _scope_data(conn, scope)
        if not imports and not notes and not bases and not fees:
            raise ReconciliationNotFound("Không tìm thấy dữ liệu đối chiếu của tài khoản đã chọn")
        _check_revision(expected_revision, _scope_revision(scope, imports, notes, bases, fees))
        deleted_notes = conn.execute("DELETE FROM reconciliation_notes WHERE scope=?", (scope,)).rowcount
        deleted_imports = conn.execute("DELETE FROM reconciliation_imports WHERE scope=?", (scope,)).rowcount
        deleted_bases = conn.execute("DELETE FROM reconciliation_bases WHERE scope=?", (scope,)).rowcount
        deleted_fees = conn.execute("DELETE FROM reconciliation_fees WHERE scope=?", (scope,)).rowcount
    return dict(scope=scope, deleted_imports_count=deleted_imports,
                deleted_notes_count=deleted_notes, deleted_bases_count=deleted_bases, deleted_fees_count=deleted_fees)


def save_fees(store, scope, buy_fee_percent, sell_fee_percent, sell_tax_percent, username):
    """Configure explicit percentage rates for all history of an existing scope."""
    if not isinstance(scope, str) or not scope:
        raise ValueError("Cần chọn tài khoản đối chiếu")
    buy_fee = _percent(buy_fee_percent, "Phí mua")
    sell_fee = _percent(sell_fee_percent, "Phí bán")
    sell_tax = _percent(sell_tax_percent, "Thuế bán")
    if sell_fee + sell_tax >= 100:
        raise ValueError("Tổng phí bán và thuế bán phải nhỏ hơn 100%")
    record = dict(scope=scope, buy_fee_percent=_exact(buy_fee), sell_fee_percent=_exact(sell_fee),
                  sell_tax_percent=_exact(sell_tax), updated_by=username, updated_at=utc_now_iso())
    with store.connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        if not _scope_exists(conn, scope):
            raise ValueError("Không tìm thấy tài khoản đối chiếu đã có dữ liệu")
        conn.execute("INSERT INTO reconciliation_fees(scope,buy_fee_percent,sell_fee_percent,sell_tax_percent,updated_by,updated_at) VALUES(?,?,?,?,?,?) "
                     "ON CONFLICT(scope) DO UPDATE SET buy_fee_percent=excluded.buy_fee_percent,sell_fee_percent=excluded.sell_fee_percent,"
                     "sell_tax_percent=excluded.sell_tax_percent,updated_by=excluded.updated_by,updated_at=excluded.updated_at",
                     (scope, record["buy_fee_percent"], record["sell_fee_percent"], record["sell_tax_percent"], username, record["updated_at"]))
    return _public_fees(scope, record)


def build_sell_plan(store, scope, ticker, parts, target_profit_percent, upper_profit_percent=None, price_step=100):
    """Suggest equal quantity parts of current surplus without placing orders."""
    if not isinstance(scope, str) or not scope:
        raise ValueError("Cần chọn tài khoản đối chiếu")
    if not isinstance(ticker, str) or not re.fullmatch(r"[A-Z0-9][A-Z0-9._-]{0,19}", ticker.strip().upper()):
        raise ValueError("Mã cổ phiếu không hợp lệ")
    ticker = ticker.strip().upper()
    if isinstance(parts, bool) or not isinstance(parts, int) or not 1 <= parts <= 100:
        raise ValueError("Số phần bán phải là số nguyên từ 1 đến 100")
    if isinstance(price_step, bool) or not isinstance(price_step, int) or not 1 <= price_step <= 1000000:
        raise ValueError("Bước giá phải là số nguyên từ 1 đến 1.000.000 VND")
    lower = _percent(target_profit_percent, "Lợi nhuận tối thiểu", maximum=Decimal(10000), inclusive=True)
    upper = lower if upper_profit_percent is None else _percent(upper_profit_percent, "Lợi nhuận tối đa", maximum=Decimal(10000), inclusive=True)
    if upper < lower:
        raise ValueError("Lợi nhuận tối đa phải lớn hơn hoặc bằng lợi nhuận tối thiểu")
    report, holdings, fee_record = _build_report_snapshot(store, scope)
    holding = holdings.get((scope, ticker))
    if not holding or not holding["quantity"]:
        raise ValueError("Mã đã chọn không có tồn mua vượt cơ sở để chia bán")
    quantity, cost = holding["quantity"], holding["cost"]
    if parts > quantity:
        raise ValueError("Số phần bán không được vượt số cổ phiếu tồn mua")
    average = cost / quantity
    sell_fee_rate = Fraction(Decimal(fee_record["sell_fee_percent"])) / 100 if fee_record else Fraction(0)
    sell_tax_rate = Fraction(Decimal(fee_record["sell_tax_percent"])) / 100 if fee_record else Fraction(0)
    lower, upper = Fraction(lower), Fraction(upper)
    per_part, remainder = divmod(quantity, parts)
    orders = []
    total = dict(quantity=0, **{key: Fraction(0) for key in ("allocated_cost", "gross_proceeds", "sell_fee", "sell_tax", "net_proceeds", "net_profit")})
    def target_price(target):
        required_steps = average * (1 + target / 100) / (1 - sell_fee_rate - sell_tax_rate) / price_step
        # Exact rational ceiling avoids adding a tick when a repeating average
        # mathematically lands on the tick boundary (e.g. 2/3 × 1.5 = 1).
        return ((required_steps.numerator + required_steps.denominator - 1) // required_steps.denominator) * price_step

    for index in range(parts):
        part_quantity = per_part + int(index < remainder)
        target = lower + (upper - lower) * index / (parts - 1) if parts > 1 else lower
        price = target_price(target)
        allocated = average * part_quantity
        gross = Fraction(price * part_quantity)
        fee, tax = gross * sell_fee_rate, gross * sell_tax_rate
        net = gross - fee - tax
        profit = net - allocated
        values = dict(quantity=part_quantity, allocated_cost=allocated, gross_proceeds=gross,
                      sell_fee=fee, sell_tax=tax, net_proceeds=net, net_profit=profit)
        orders.append(dict(part=index + 1, target_profit_percent=float(target), price=price,
                           **{key: int(value) if key == "quantity" else float(value) for key, value in values.items()},
                           net_profit_percent=float(profit / allocated * 100)))
        for key, value in values.items():
            total[key] += value
    return dict(scope=scope, ticker=ticker, quantity=quantity, average_cost=float(average), total_cost=float(cost),
                fees=report["fees"], target_profit_percent=float(lower), upper_profit_percent=float(upper), price_step=price_step,
                price_min=target_price(lower), price_max=target_price(upper), orders=orders,
                totals={key: int(value) if key == "quantity" else float(value) for key, value in total.items()} |
                       dict(net_profit_percent=float(total["net_profit"] / cost * 100)),
                revision=report["revision"], currency="VND", cost_includes_buy_fee=True)
